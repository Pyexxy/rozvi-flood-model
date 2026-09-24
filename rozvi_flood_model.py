#!/usr/bin/env python3
"""
Rozvi Flood Model
=================

Corridor-scale flood susceptibility screening for linear assets (roads).

The model scores fixed-length road segments using a weighted combination of
terrain and hydro-meteorological drivers, maps scores onto a 0-100 risk
category scale, and writes tabular, geospatial and Word report outputs.

Scope
-----
This is a screening tool. It does not compute return-period flood depths,
defence performance, climate-adjusted depth surfaces, or financial loss
(EAL/PML). Climate information (CMIP6 / emerging CMIP7 pathways) is reported
as regional context only. Those limitations are stated explicitly in the
outputs so results are not misread as calibrated hazard or loss products.

Usage
-----
    python rozvi_flood_model.py --roads data/roads.geojson --outdir outputs

    python rozvi_flood_model.py \\
        --roads data/roads.geojson \\
        --dem data/dem.tif \\
        --outdir outputs \\
        --portfolio-name "Harare-Beitbridge Road"

Dependencies are listed in requirements.txt.
"""

from __future__ import annotations

import argparse
import json
import math
import sys
import uuid
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd

try:
    import geopandas as gpd
    from shapely.geometry import LineString, MultiLineString, mapping
    from shapely.ops import substring
except ImportError as exc:
    sys.exit("geopandas and shapely are required. Install from requirements.txt\n" + str(exc))

try:
    import rasterio
    from rasterio import features, transform as rio_transform
    from rasterio.enums import Resampling
    from rasterio.warp import calculate_default_transform, reproject
    from rasterio.windows import from_bounds
except ImportError as exc:
    sys.exit("rasterio is required. Install from requirements.txt\n" + str(exc))

try:
    from scipy import ndimage
except ImportError:
    ndimage = None

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch
    HAS_MPL = True
except ImportError:
    HAS_MPL = False


# ---------------------------------------------------------------------------
# Model constants
# ---------------------------------------------------------------------------

MODEL_NAME = "Rozvi Flood Model"
MODEL_VERSION = "1.3.0"

# Driver weights (must sum to 1.0). Applied after each driver is scaled to 1-10.
DEFAULT_WEIGHTS = {
    "precip": 0.30,
    "dist_water": 0.25,
    "lulc": 0.20,
    "slope": 0.15,
    "dem": 0.10,
}

# Human-readable driver catalogue for reports and audit trails
DRIVER_CATALOGUE = [
    {
        "key": "precip",
        "name": "Precipitation",
        "weight": DEFAULT_WEIGHTS["precip"],
        "direction": "Higher rainfall increases susceptibility",
        "data_status": "Local rainfall raster if supplied; otherwise uniform placeholder",
        "scaling": "AOI percentile stretch (2nd-98th); not inverted",
    },
    {
        "key": "dist_water",
        "name": "Proximity to water / local low points",
        "weight": DEFAULT_WEIGHTS["dist_water"],
        "direction": "Greater distance (or height above local minima) reduces susceptibility",
        "data_status": "Water mask distance if supplied; otherwise height above focal DEM minimum",
        "scaling": "AOI percentile stretch; inverted",
    },
    {
        "key": "lulc",
        "name": "Land cover",
        "weight": DEFAULT_WEIGHTS["lulc"],
        "direction": "Class-dependent when a land-cover raster is available",
        "data_status": "Neutral value (5) when no land-cover raster is supplied",
        "scaling": "Fixed class map or neutral default",
    },
    {
        "key": "slope",
        "name": "Terrain slope",
        "weight": DEFAULT_WEIGHTS["slope"],
        "direction": "Steeper slopes reduce ponding susceptibility in this formulation",
        "data_status": "Derived from DEM",
        "scaling": "AOI percentile stretch; inverted",
    },
    {
        "key": "dem",
        "name": "Elevation",
        "weight": DEFAULT_WEIGHTS["dem"],
        "direction": "Higher elevation reduces susceptibility",
        "data_status": "DEM GeoTIFF, optional public SRTM, or synthetic test surface",
        "scaling": "AOI percentile stretch; inverted",
    },
]

SCREENING_THRESHOLD = 7.0
SEGMENT_LENGTH_M = 100.0
CORRIDOR_HALF_WIDTH_M = 15.0
MAX_SEGMENTS = 5000

RISK_BANDS = [
    ("Minimal", 0, 10),
    ("Low", 11, 30),
    ("Moderate", 31, 60),
    ("High", 61, 85),
    ("Extreme", 86, 100),
]

BAND_COLORS = {
    "Minimal": "#2ca02c",
    "Low": "#98df8a",
    "Moderate": "#ffbb78",
    "High": "#ff7f0e",
    "Extreme": "#d62728",
}

# Climate context (documentary only in this build)
CLIMATE_CONTEXT = {
    "framework": "CMIP6 (CMIP7 not yet used operationally in this screening build)",
    "pathways_referenced": "SSP1-2.6, SSP2-4.5, SSP3-7.0, SSP5-8.5 (for narrative context)",
    "application": (
        "Regional seasonal anomaly context only. No segment-scale depth or "
        "probability adjustment is applied from CMIP outputs in this version."
    ),
    "cmip7_note": (
        "CMIP7 is the next generation of the Coupled Model Intercomparison Project. "
        "As of this model version, screening reports reference CMIP6 pathways for "
        "context. CMIP7-based regional products may be adopted in a later release "
        "once stable, peer-reviewed downscaled datasets are available for the study region."
    ),
    "segment_scale_climate": "not assessed",
}

# Model quality / assurance statements for reports
MODEL_QUALITY = {
    "purpose": "Relative susceptibility ranking for corridor screening and prioritisation",
    "calibration_status": "Not calibrated against observed road closures or gauged flood peaks",
    "validation_status": (
        "No independent closure or inundation validation in the default workflow. "
        "Users should document any project-specific checks separately."
    ),
    "spatial_unit": "Fixed-length road segments with corridor buffer sampling",
    "vertical_data_limit": (
        "DEM resolution (e.g. SRTM ~30 m) cannot resolve embankments or structures "
        "below approximately 2-3 m."
    ),
    "missing_data_policy": "Unsampled or missing driver pixels contribute a neutral score of 5 before weighting",
    "suitable_uses": "Prioritisation of inspection, further study, and data collection",
    "unsuitable_uses": (
        "Design water levels, formal flood zoning, insurance pricing, or statements of "
        "absolute flood probability without additional calibrated modelling"
    ),
}


def risk_category_legacy(score: Optional[float]) -> str:
    """Four-class label retained for continuity with earlier corridor reports."""
    if score is None or (isinstance(score, float) and math.isnan(score)):
        return "N/A"
    if score <= 3:
        return "LOW"
    if score <= 6:
        return "MODERATE"
    if score <= 8:
        return "HIGH"
    return "VERY HIGH / CRITICAL"


def score_to_rozvi_category(score: float) -> int:
    """Map susceptibility score (1-10) to risk category (0-100)."""
    s = float(np.clip(score, 1.0, 10.0))
    return int(round((s - 1.0) / 9.0 * 100.0))


def rozvi_band(category: int) -> str:
    """Return band name for a 0-100 category value."""
    for name, lo, hi in RISK_BANDS:
        if lo <= category <= hi:
            return name
    return "Minimal" if category < 0 else "Extreme"


def relative_risk_score_from_category(category: int) -> int:
    """
    Relative risk on a 0-1,000,000 scale.

    Implemented as category * 10,000 for local ranking within a run.
    Not normalised to a national or global exposure baseline.
    """
    return int(np.clip(category, 0, 100)) * 10_000


def weights_sum_ok(weights: Dict[str, float], tol: float = 1e-6) -> bool:
    return abs(sum(weights.values()) - 1.0) <= tol


# ---------------------------------------------------------------------------
# Geometry
# ---------------------------------------------------------------------------

def ensure_wgs84(gdf: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    if gdf.crs is None:
        return gdf.set_crs(epsg=4326)
    if gdf.crs.to_epsg() != 4326:
        return gdf.to_crs(epsg=4326)
    return gdf


def to_metric_crs(gdf: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    """Project to a local UTM zone for length and buffer operations in metres."""
    gdf = ensure_wgs84(gdf)
    try:
        cen = gdf.unary_union.centroid
        zone = int((cen.x + 180) // 6) + 1
        epsg = 32600 + zone if cen.y >= 0 else 32700 + zone
        return gdf.to_crs(epsg=epsg)
    except Exception:
        return gdf.to_crs(epsg=3857)


def cut_line_to_segments(geom, seg_len_m: float) -> List[LineString]:
    """Split a line into pieces of approximately seg_len_m (metric CRS)."""
    if geom is None or geom.is_empty:
        return []
    if isinstance(geom, MultiLineString):
        parts: List[LineString] = []
        for g in geom.geoms:
            parts.extend(cut_line_to_segments(g, seg_len_m))
        return parts
    if not isinstance(geom, LineString):
        return []
    length = geom.length
    if length == 0:
        return []
    n = max(1, int(math.ceil(length / seg_len_m)))
    segs = []
    for i in range(n):
        start = i * seg_len_m
        end = min((i + 1) * seg_len_m, length)
        if end - start < 1e-6:
            continue
        try:
            piece = substring(geom, start, end)
            if piece is not None and not piece.is_empty:
                segs.append(piece)
        except Exception:
            continue
    return segs


def segment_roads(
    roads_gdf: gpd.GeoDataFrame,
    seg_len_m: float = SEGMENT_LENGTH_M,
    max_segments: int = MAX_SEGMENTS,
) -> gpd.GeoDataFrame:
    """
    Explode a road network into fixed-length segments.

    Returns a GeoDataFrame in EPSG:4326 with columns
    road_id, seg_index, geometry, length_m.
    """
    metric = to_metric_crs(roads_gdf)
    rows = []
    for idx, row in metric.iterrows():
        pieces = cut_line_to_segments(row.geometry, seg_len_m)
        for j, piece in enumerate(pieces):
            rows.append(
                {
                    "road_id": str(idx),
                    "seg_index": j,
                    "geometry": piece,
                    "length_m": piece.length,
                }
            )
            if len(rows) >= max_segments:
                break
        if len(rows) >= max_segments:
            break
    if not rows:
        raise RuntimeError(
            "No line geometries found. Provide LineString or MultiLineString features."
        )
    return gpd.GeoDataFrame(rows, crs=metric.crs).to_crs(epsg=4326)


# ---------------------------------------------------------------------------
# Raster helpers
# ---------------------------------------------------------------------------

def read_raster_window(
    path: Path,
    bounds: Tuple[float, float, float, float],
) -> Tuple[np.ndarray, rasterio.Affine]:
    """Read a WGS84 window from a GeoTIFF; reproject if the source CRS differs."""
    with rasterio.open(path) as src:
        from pyproj import Transformer

        dst_crs = "EPSG:4326"
        if src.crs and src.crs.to_string() != dst_crs:
            tf = Transformer.from_crs(dst_crs, src.crs, always_xy=True)
            x0, y0 = tf.transform(bounds[0], bounds[1])
            x1, y1 = tf.transform(bounds[2], bounds[3])
            rb = (min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1))
        else:
            rb = bounds

        window = from_bounds(*rb, transform=src.transform)
        data = src.read(1, window=window, boundless=True, fill_value=src.nodata or np.nan)
        win_transform = src.window_transform(window)

        if src.crs and src.crs.to_string() != dst_crs:
            h, w = data.shape
            dst_transform, dw, dh = calculate_default_transform(
                src.crs, dst_crs, w, h, *rb
            )
            dst = np.full((dh, dw), np.nan, dtype=np.float64)
            reproject(
                source=data.astype(np.float64),
                destination=dst,
                src_transform=win_transform,
                src_crs=src.crs,
                dst_transform=dst_transform,
                dst_crs=dst_crs,
                resampling=Resampling.bilinear,
                src_nodata=src.nodata,
                dst_nodata=np.nan,
            )
            return dst, dst_transform
        return data.astype(np.float64), win_transform


def slope_from_dem(dem: np.ndarray, transform: rasterio.Affine) -> np.ndarray:
    """Slope in degrees from a DEM array."""
    px = abs(transform.a)
    py = abs(transform.e)
    if px < 0.1:
        lat_m = 111_320.0
        lon_m = 111_320.0 * math.cos(math.radians(20))
        dx, dy = px * lon_m, py * lat_m
    else:
        dx, dy = px, py
    gy, gx = np.gradient(dem.astype(np.float64), dy, dx)
    return np.degrees(np.arctan(np.sqrt(gx ** 2 + gy ** 2)))


def zonal_mean(geom_wgs, array: np.ndarray, transform: rasterio.Affine) -> Optional[float]:
    """Mean of raster values under a geometry. Returns None if no valid pixels."""
    try:
        mask = features.geometry_mask(
            [mapping(geom_wgs)],
            out_shape=array.shape,
            transform=transform,
            invert=True,
        )
        vals = array[mask]
        vals = vals[np.isfinite(vals)]
        if vals.size == 0:
            return None
        return float(np.mean(vals))
    except Exception:
        return None


def adaptive_unit_scale(
    values: np.ndarray,
    low_pct: float = 2,
    high_pct: float = 98,
    invert: bool = False,
) -> np.ndarray:
    """
    Scale an array to 1-10 using AOI percentiles.

    invert=True assigns lower risk to higher original values (e.g. elevation).
    Missing values are mapped to a neutral score of 5.
    """
    finite = values[np.isfinite(values)]
    if finite.size < 5:
        return np.full_like(values, 5.0, dtype=np.float64)
    lo = np.nanpercentile(finite, low_pct)
    hi = np.nanpercentile(finite, high_pct)
    if hi <= lo:
        hi = lo + 1.0
    scaled = np.clip((values - lo) / (hi - lo), 0, 1)
    risk = (1.0 - scaled) * 9.0 + 1.0 if invert else scaled * 9.0 + 1.0
    risk = np.where(np.isfinite(values), risk, 5.0)
    return np.clip(risk, 1.0, 10.0)


def try_download_srtm(bounds: Tuple[float, float, float, float], dest: Path) -> Optional[Path]:
    """Optional public SRTM clip via the elevation package."""
    try:
        import elevation  # type: ignore
    except ImportError:
        print("  [info] Optional package 'elevation' not installed; skipping SRTM download.")
        return None
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    minx, miny, maxx, maxy = bounds
    margin = 0.05
    bb = (minx - margin, miny - margin, maxx + margin, maxy + margin)
    try:
        print(f"  Downloading SRTM for bbox {bb} ...")
        elevation.clip(bounds=bb, output=str(dest), product="SRTM3")
        elevation.clean()
        if dest.exists():
            return dest
    except Exception as exc:
        print(f"  [warn] SRTM download failed: {exc}")
    return None


def make_synthetic_dem(
    bounds: Tuple[float, float, float, float],
    shape: Tuple[int, int] = (400, 400),
) -> Tuple[np.ndarray, rasterio.Affine]:
    """
    Deterministic synthetic DEM for pipeline tests when no DEM is available.

    Not suitable for operational screening.
    """
    minx, miny, maxx, maxy = bounds
    h, w = shape
    xs = np.linspace(minx, maxx, w)
    ys = np.linspace(maxy, miny, h)
    xx, yy = np.meshgrid(xs, ys)
    dem = 500 + 80 * np.sin((xx - minx) * 3) + 40 * np.cos((yy - miny) * 2)
    dem += np.random.default_rng(42).normal(0, 5, size=dem.shape)
    transform = rio_transform.from_bounds(minx, miny, maxx, maxy, w, h)
    return dem.astype(np.float64), transform


# ---------------------------------------------------------------------------
# Core screening
# ---------------------------------------------------------------------------

def build_risk_layers(
    bounds: Tuple[float, float, float, float],
    dem_path: Optional[Path] = None,
    water_path: Optional[Path] = None,
    rain_path: Optional[Path] = None,
    weights: Optional[Dict[str, float]] = None,
) -> Dict[str, Any]:
    """Build driver rasters and the composite susceptibility surface over bounds."""
    w = dict(weights or DEFAULT_WEIGHTS)
    if not weights_sum_ok(w):
        print(f"  [warn] Driver weights sum to {sum(w.values()):.4f}, expected 1.0")

    dem_source = "synthetic"
    rain_source = "placeholder (80 mm uniform)"
    water_source = "HAND-like focal minimum proxy"

    if dem_path and Path(dem_path).exists():
        print(f"  Reading DEM: {dem_path}")
        dem_arr, transform = read_raster_window(Path(dem_path), bounds)
        dem_source = str(dem_path)
    else:
        srtm = try_download_srtm(bounds, Path("cache_srtm_clip.tif"))
        if srtm and srtm.exists():
            dem_arr, transform = read_raster_window(srtm, bounds)
            dem_source = "SRTM3 (elevation package)"
        else:
            print("  [demo] No DEM supplied; using synthetic surface for testing only.")
            dem_arr, transform = make_synthetic_dem(bounds)
            dem_source = "synthetic (demo)"

    dem_arr = np.where((dem_arr < -100) | (dem_arr > 9000), np.nan, dem_arr)
    slope_arr = slope_from_dem(dem_arr, transform)

    if water_path and Path(water_path).exists():
        water, _ = read_raster_window(Path(water_path), bounds)
        from scipy.ndimage import distance_transform_edt

        inv = (water <= 0).astype(np.uint8)
        dist_px = distance_transform_edt(inv)
        px = abs(transform.a)
        scale = px * 111_320 if px < 0.1 else px
        dist_arr = dist_px * scale
        water_source = str(water_path)
    elif ndimage is not None:
        focal_min = ndimage.minimum_filter(
            np.nan_to_num(dem_arr, nan=np.nanmax(dem_arr)), size=15
        )
        dist_arr = np.clip(dem_arr - focal_min, 0, None)
    else:
        dist_arr = np.zeros_like(dem_arr)

    if rain_path and Path(rain_path).exists():
        rain_arr, _ = read_raster_window(Path(rain_path), bounds)
        rain_source = str(rain_path)
    else:
        rain_arr = np.full_like(dem_arr, 80.0)

    dem_r = adaptive_unit_scale(dem_arr, invert=True)
    slope_r = adaptive_unit_scale(slope_arr, invert=True)
    dist_r = adaptive_unit_scale(dist_arr, invert=True)
    precip_r = adaptive_unit_scale(rain_arr, invert=False)
    lulc_r = np.full_like(dem_arr, 5.0)

    risk = (
        precip_r * w["precip"]
        + dist_r * w["dist_water"]
        + lulc_r * w["lulc"]
        + slope_r * w["slope"]
        + dem_r * w["dem"]
    )
    risk = np.clip(risk, 1.0, 10.0)

    # Refresh catalogue data_status with actual sources for this run
    catalogue = []
    for d in DRIVER_CATALOGUE:
        entry = dict(d)
        entry["weight"] = w[d["key"]]
        if d["key"] == "precip":
            entry["data_status"] = rain_source
        elif d["key"] == "dist_water":
            entry["data_status"] = water_source
        elif d["key"] == "dem":
            entry["data_status"] = dem_source
        elif d["key"] == "slope":
            entry["data_status"] = f"Derived from DEM ({dem_source})"
        catalogue.append(entry)

    return {
        "risk": risk,
        "dem_arr": dem_arr,
        "slope_arr": slope_arr,
        "dist_arr": dist_arr,
        "rain_arr": rain_arr,
        # Scaled driver surfaces (1-10) used for score and contribution
        "precip_r": precip_r,
        "dist_r": dist_r,
        "lulc_r": lulc_r,
        "slope_r": slope_r,
        "dem_r": dem_r,
        "transform": transform,
        "bounds": bounds,
        "dem_source": dem_source,
        "rain_source": rain_source,
        "water_source": water_source,
        "weights": w,
        "driver_catalogue": catalogue,
    }


def score_segments(
    segments: gpd.GeoDataFrame,
    layers: Dict[str, Any],
    corridor_half_width_m: float = CORRIDOR_HALF_WIDTH_M,
) -> List[Dict[str, Any]]:
    """
    Sample drivers under each segment corridor and compute scores.

    The composite score still uses the fixed global weights. For each segment,
    driver contributions are also stored so the leading factors at that location
    can be reported (contribution_d = weight_d * scaled_driver_d).
    """
    metric_segs = to_metric_crs(segments)
    transform = layers["transform"]
    w = layers.get("weights", DEFAULT_WEIGHTS)
    scored: List[Dict[str, Any]] = []

    driver_keys = ("precip", "dist_water", "lulc", "slope", "dem")
    scaled_layers = {
        "precip": layers["precip_r"],
        "dist_water": layers["dist_r"],
        "lulc": layers["lulc_r"],
        "slope": layers["slope_r"],
        "dem": layers["dem_r"],
    }
    name_map = {
        "precip": "Precipitation",
        "dist_water": "Water proximity / local low",
        "lulc": "Land cover",
        "slope": "Slope",
        "dem": "Elevation",
    }

    for i, (_, row) in enumerate(segments.iterrows()):
        geom = row.geometry
        mrow = metric_segs.iloc[i]
        buf = mrow.geometry.buffer(corridor_half_width_m)
        buf_wgs = gpd.GeoSeries([buf], crs=metric_segs.crs).to_crs(epsg=4326).iloc[0]

        dem_m = zonal_mean(buf_wgs, layers["dem_arr"], transform)
        slope_m = zonal_mean(buf_wgs, layers["slope_arr"], transform)
        dist_m = zonal_mean(buf_wgs, layers["dist_arr"], transform)
        rain_m = zonal_mean(buf_wgs, layers["rain_arr"], transform)

        # Scaled driver values (1-10) and weighted contributions at this segment
        scaled: Dict[str, float] = {}
        contrib: Dict[str, float] = {}
        for key in driver_keys:
            val = zonal_mean(buf_wgs, scaled_layers[key], transform)
            if val is None:
                val = 5.0
            scaled[key] = float(val)
            contrib[key] = float(w[key] * scaled[key])

        final = float(np.clip(sum(contrib.values()), 1.0, 10.0))
        # Rank drivers by contribution (highest first) for this location
        ranked = sorted(contrib.items(), key=lambda kv: kv[1], reverse=True)
        top_drivers = "; ".join(
            f"{name_map[k]} ({c:.2f})" for k, c in ranked[:3]
        )
        primary_driver = name_map[ranked[0][0]] if ranked else "N/A"

        cat = score_to_rozvi_category(final)
        band = rozvi_band(cat)
        cen = geom.centroid if geom and not geom.is_empty else None

        scored.append(
            {
                "uid": f"R-{i + 1:04d}",
                "road_id": row.get("road_id", i),
                "seg_index": int(row.get("seg_index", i)),
                "geometry": geom,
                "base_risk": round(final, 2),
                "road_score": round(final, 2),
                "final_score": round(final, 2),
                "category": risk_category_legacy(final),
                "risk_category_0_100": cat,
                "band": band,
                "relative_risk_score": relative_risk_score_from_category(cat),
                "dominant_peril": "pluvial",
                # Physical samples
                "elev_m": None if dem_m is None else round(dem_m, 1),
                "slope_deg": None if slope_m is None else round(slope_m, 2),
                "dist_water_proxy_m": None if dist_m is None else round(dist_m, 1),
                "forecast_rain_mm": None if rain_m is None else round(rain_m, 1),
                # Scaled drivers (1-10) at this segment
                "drv_precip_scaled": round(scaled["precip"], 2),
                "drv_dist_water_scaled": round(scaled["dist_water"], 2),
                "drv_lulc_scaled": round(scaled["lulc"], 2),
                "drv_slope_scaled": round(scaled["slope"], 2),
                "drv_dem_scaled": round(scaled["dem"], 2),
                # Weighted contributions (weight * scaled); sum ~= final_score
                "contrib_precip": round(contrib["precip"], 3),
                "contrib_dist_water": round(contrib["dist_water"], 3),
                "contrib_lulc": round(contrib["lulc"], 3),
                "contrib_slope": round(contrib["slope"], 3),
                "contrib_dem": round(contrib["dem"], 3),
                "primary_driver": primary_driver,
                "top_drivers": top_drivers,
                "lon": None if cen is None else round(cen.x, 6),
                "lat": None if cen is None else round(cen.y, 6),
                "length_m": float(row.get("length_m", SEGMENT_LENGTH_M)),
                "depth_100_undef_m": None,
                "eal_present": None,
                "eal_future": None,
            }
        )
    return scored


def assign_chainage(scored: List[Dict], seg_len: float) -> List[Dict]:
    """Assign sequential chainage and section IDs per parent road feature."""
    by_road: Dict[str, List] = defaultdict(list)
    for s in scored:
        by_road[str(s["road_id"])].append(s)
    for rid, segs in by_road.items():
        segs.sort(key=lambda x: x["seg_index"])
        chain = 0.0
        for i, s in enumerate(segs):
            s["chainage_start_m"] = round(chain, 1)
            s["chainage_end_m"] = round(chain + seg_len, 1)
            s["chainage_mid_m"] = round(chain + seg_len / 2.0, 1)
            s["section_id"] = f"{rid}-{i + 1:03d}"
            s["asset_id"] = s["section_id"]
            s["asset_name"] = f"Road segment {s['section_id']}"
            s["asset_type"] = "road_segment"
            chain = s["chainage_end_m"]
    ranked = sorted(scored, key=lambda x: x["risk_category_0_100"], reverse=True)
    for rank, s in enumerate(ranked, start=1):
        s["rank"] = rank
    return scored


def portfolio_metrics(scored: List[Dict], seg_len: float) -> Dict[str, Any]:
    n = len(scored)
    total_km = n * seg_len / 1000.0
    exposed = sum(1 for s in scored if s["final_score"] >= SCREENING_THRESHOLD)
    band_counts = {name: 0 for name, _, _ in RISK_BANDS}
    for s in scored:
        band_counts[s["band"]] = band_counts.get(s["band"], 0) + 1
    cats = [s["risk_category_0_100"] for s in scored]
    top = max(scored, key=lambda x: x["risk_category_0_100"]) if scored else None
    return {
        "asset_count": n,
        "road_length_km": round(total_km, 2),
        "n_assets_exposed_screening": exposed,
        "pct_assets_exposed_screening": round(exposed / n * 100, 1) if n else 0,
        "portfolio_mean_category": round(float(np.mean(cats)), 1) if cats else 0.0,
        "top_asset_name": top["asset_name"] if top else None,
        "top_asset_id": top["asset_id"] if top else None,
        "top_asset_category_score": top["risk_category_0_100"] if top else None,
        "band_counts": band_counts,
        "dominant_peril": "pluvial",
    }



# ---------------------------------------------------------------------------
# Figures (maps and charts for the Word report)
# ---------------------------------------------------------------------------

def _sorted_along_network(scored: List[Dict]) -> List[Dict]:
    """Order segments by road then chainage for longitudinal plots."""
    return sorted(
        scored,
        key=lambda s: (str(s.get("road_id", "")), float(s.get("chainage_mid_m") or 0.0)),
    )


def generate_figures(
    scored: List[Dict],
    metrics: Dict[str, Any],
    out_dir: Path,
) -> Dict[str, Path]:
    """
    Build insight figures from scored segments.

    Returns a dict of figure keys to PNG paths. Requires matplotlib.
    """
    paths: Dict[str, Path] = {}
    if not HAS_MPL or not scored:
        if not HAS_MPL:
            print("  [info] matplotlib not installed; skipping figures.")
        return paths

    fig_dir = out_dir / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    n = len(scored)
    bc = metrics.get("band_counts", {})

    # --- 1. Risk band distribution (km and count) ---
    try:
        fig, ax = plt.subplots(figsize=(7.2, 3.6))
        names = [b[0] for b in RISK_BANDS]
        counts = [bc.get(name, 0) for name in names]
        colors = [BAND_COLORS[name] for name in names]
        bars = ax.barh(names, counts, color=colors, edgecolor="white")
        ax.set_xlabel("Number of segments")
        ax.set_title("Risk band distribution")
        for bar, c in zip(bars, counts):
            if c > 0:
                ax.text(
                    bar.get_width() + max(counts) * 0.01,
                    bar.get_y() + bar.get_height() / 2,
                    str(c),
                    va="center",
                    fontsize=8,
                )
        ax.set_xlim(0, max(counts) * 1.15 if max(counts) else 1)
        fig.tight_layout()
        path = fig_dir / "band_distribution.png"
        fig.savefig(path, dpi=150)
        plt.close(fig)
        paths["band_distribution"] = path
    except Exception as exc:
        print(f"  [warn] band chart failed: {exc}")

    # --- 2. Score histogram ---
    try:
        fig, ax = plt.subplots(figsize=(7.2, 3.4))
        scores = [float(s["final_score"]) for s in scored]
        ax.hist(scores, bins=18, range=(1, 10), color="#4c78a8", edgecolor="white")
        ax.axvline(SCREENING_THRESHOLD, color="#d62728", linestyle="--", linewidth=1.2, label=f"Screening threshold ({SCREENING_THRESHOLD})")
        ax.set_xlabel("Susceptibility score (1–10)")
        ax.set_ylabel("Segment count")
        ax.set_title("Distribution of susceptibility scores")
        ax.legend(fontsize=8)
        fig.tight_layout()
        path = fig_dir / "score_histogram.png"
        fig.savefig(path, dpi=150)
        plt.close(fig)
        paths["score_histogram"] = path
    except Exception as exc:
        print(f"  [warn] histogram failed: {exc}")

    # --- 3. Longitudinal profile along network order ---
    try:
        ordered = _sorted_along_network(scored)
        # cumulative distance index (m) so multi-road corridors plot as one sequence
        x = []
        y = []
        colors = []
        cursor = 0.0
        for s in ordered:
            length = float(s.get("length_m") or SEGMENT_LENGTH_M)
            cursor += length
            x.append(cursor / 1000.0)  # km
            y.append(float(s["final_score"]))
            colors.append(BAND_COLORS.get(s.get("band"), "#999999"))
        fig, ax = plt.subplots(figsize=(8.5, 3.6))
        ax.scatter(x, y, c=colors, s=8, alpha=0.85, linewidths=0)
        ax.plot(x, y, color="#333333", linewidth=0.4, alpha=0.35)
        ax.axhline(SCREENING_THRESHOLD, color="#d62728", linestyle="--", linewidth=1.0)
        ax.set_xlabel("Aligned distance along assessed network (km)")
        ax.set_ylabel("Susceptibility score")
        ax.set_ylim(1, 10)
        ax.set_title("Longitudinal risk profile")
        legend_items = [Patch(facecolor=BAND_COLORS[n], label=n) for n in names]
        ax.legend(handles=legend_items, fontsize=7, loc="upper right", ncol=3)
        fig.tight_layout()
        path = fig_dir / "longitudinal_profile.png"
        fig.savefig(path, dpi=150)
        plt.close(fig)
        paths["longitudinal_profile"] = path
    except Exception as exc:
        print(f"  [warn] longitudinal profile failed: {exc}")

    # --- 4. Driver contributions for top segments ---
    try:
        top = sorted(scored, key=lambda s: s["risk_category_0_100"], reverse=True)[:12]
        labels = [str(s.get("section_id")) for s in top]
        keys = ["contrib_precip", "contrib_dist_water", "contrib_lulc", "contrib_slope", "contrib_dem"]
        key_labels = ["Precipitation", "Water / low points", "Land cover", "Slope", "Elevation"]
        key_colors = ["#4c78a8", "#72b7b2", "#54a24b", "#eeca3b", "#f58518"]
        fig, ax = plt.subplots(figsize=(8.5, 4.2))
        bottoms = [0.0] * len(top)
        for key, lab, col in zip(keys, key_labels, key_colors):
            vals = [float(s.get(key) or 0.0) for s in top]
            ax.barh(labels, vals, left=bottoms, color=col, edgecolor="white", label=lab)
            bottoms = [b + v for b, v in zip(bottoms, vals)]
        ax.invert_yaxis()
        ax.set_xlabel("Weighted contribution to score")
        ax.set_title("Driver contributions — highest-ranked segments")
        ax.legend(fontsize=7, loc="lower right")
        fig.tight_layout()
        path = fig_dir / "driver_contributions_top.png"
        fig.savefig(path, dpi=150)
        plt.close(fig)
        paths["driver_contributions"] = path
    except Exception as exc:
        print(f"  [warn] driver contribution chart failed: {exc}")

    # --- 5. Corridor map (segments coloured by band) ---
    try:
        fig, ax = plt.subplots(figsize=(8.0, 6.0))
        plotted = 0
        for s in scored:
            geom = s.get("geometry")
            if geom is None or geom.is_empty:
                continue
            color = BAND_COLORS.get(s.get("band"), "#999999")
            if geom.geom_type == "LineString":
                xs, ys = geom.xy
                ax.plot(xs, ys, color=color, linewidth=1.6, solid_capstyle="round")
                plotted += 1
            elif geom.geom_type == "MultiLineString":
                for part in geom.geoms:
                    xs, ys = part.xy
                    ax.plot(xs, ys, color=color, linewidth=1.6, solid_capstyle="round")
                    plotted += 1
        if plotted:
            ax.set_aspect("equal", adjustable="datalim")
            ax.set_xlabel("Longitude")
            ax.set_ylabel("Latitude")
            ax.set_title("Corridor susceptibility by risk band")
            legend_items = [Patch(facecolor=BAND_COLORS[n], label=n) for n in names]
            ax.legend(handles=legend_items, fontsize=8, loc="best")
            ax.grid(True, linewidth=0.3, alpha=0.4)
            fig.tight_layout()
            path = fig_dir / "corridor_risk_map.png"
            fig.savefig(path, dpi=150)
            plt.close(fig)
            paths["corridor_map"] = path
        else:
            plt.close(fig)
    except Exception as exc:
        print(f"  [warn] corridor map failed: {exc}")

    # --- 6. Primary driver share (portfolio) ---
    try:
        from collections import Counter
        counts = Counter(s.get("primary_driver") or "N/A" for s in scored)
        labels = list(counts.keys())
        vals = [counts[k] for k in labels]
        fig, ax = plt.subplots(figsize=(7.2, 3.6))
        ax.barh(labels, vals, color="#6b8cae", edgecolor="white")
        ax.set_xlabel("Number of segments")
        ax.set_title("Primary driver (largest contribution) by segment count")
        fig.tight_layout()
        path = fig_dir / "primary_driver_share.png"
        fig.savefig(path, dpi=150)
        plt.close(fig)
        paths["primary_driver_share"] = path
    except Exception as exc:
        print(f"  [warn] primary driver chart failed: {exc}")

    print(f"  Wrote {len(paths)} figure(s) under {fig_dir}")
    return paths



# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def export_driver_register(layers: Dict[str, Any], out_dir: Path) -> Path:
    """Table of risk drivers, weights and data sources for the run."""
    rows = []
    for d in layers.get("driver_catalogue", DRIVER_CATALOGUE):
        rows.append(
            {
                "driver_key": d["key"],
                "driver_name": d["name"],
                "weight": d["weight"],
                "weight_pct": round(d["weight"] * 100, 1),
                "direction": d["direction"],
                "scaling": d["scaling"],
                "data_status": d["data_status"],
            }
        )
    path = out_dir / f"Risk_Driver_Weights_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    print(f"  Wrote {path}")
    return path


def export_executive_summary(
    scored: List[Dict],
    seg_len: float,
    metrics: Dict[str, Any],
    layers: Dict[str, Any],
    out_dir: Path,
) -> Path:
    top = sorted(scored, key=lambda x: x["final_score"], reverse=True)[:5]
    top_ids = "; ".join(s["section_id"] for s in top)
    n = metrics["asset_count"]
    total_km = metrics["road_length_km"]
    exposed_km = metrics["n_assets_exposed_screening"] * seg_len / 1000.0
    w = layers.get("weights", DEFAULT_WEIGHTS)

    row = {
        "report_title": f"{MODEL_NAME} – Executive Summary",
        "report_id": f"RFM-{datetime.now().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}",
        "report_type": "Portfolio (road segments as assets)",
        "generated_on": datetime.now().isoformat(timespec="seconds"),
        "model_version": MODEL_VERSION,
        "asset_count": n,
        "road_length_assessed_km": total_km,
        "baseline_exposed_length_km": round(exposed_km, 1),
        "baseline_exposed_pct": metrics["pct_assets_exposed_screening"],
        "portfolio_mean_category": metrics["portfolio_mean_category"],
        "top_asset_name": metrics["top_asset_name"],
        "top_asset_category_score": metrics["top_asset_category_score"],
        "dominant_peril": metrics["dominant_peril"],
        "priority_locations": top_ids,
        "weight_precip": w.get("precip"),
        "weight_dist_water": w.get("dist_water"),
        "weight_lulc": w.get("lulc"),
        "weight_slope": w.get("slope"),
        "weight_dem": w.get("dem"),
        "climate_framework": CLIMATE_CONTEXT["framework"],
        "climate_application": CLIMATE_CONTEXT["application"],
        "model_calibration_status": MODEL_QUALITY["calibration_status"],
        "model_validation_status": MODEL_QUALITY["validation_status"],
        "exposure_definition": (
            f"Segments with susceptibility score >= {SCREENING_THRESHOLD}. "
            "Risk category 0-100 is a linear map of the 1-10 score. "
            "Not a calibrated return-period probability or depth product."
        ),
        "plain_language_summary": (
            f"Approximately {total_km:.1f} km ({n} segments) assessed. "
            f"{metrics['pct_assets_exposed_screening']}% meet the screening threshold. "
            f"Mean risk category = {metrics['portfolio_mean_category']}/100. "
            f"Highest segment: {metrics['top_asset_name']} "
            f"(category {metrics['top_asset_category_score']}). "
            f"Priority sections: {top_ids}."
        ),
        "important_caveats": (
            "Screening ranking only. No AEP depth grids, defence scenarios, "
            "segment-scale climate depth adjustment, or EAL/PML in this build."
        ),
    }
    path = out_dir / f"Executive_Summary_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    pd.DataFrame([row]).to_csv(path, index=False)
    print(f"  Wrote {path}")
    return path


def export_section_register(scored: List[Dict], seg_len: float, out_dir: Path) -> Path:
    rows = []
    for s in scored:
        rows.append(
            {
                "section_id": s.get("section_id"),
                "asset_id": s.get("asset_id"),
                "road_id": s["road_id"],
                "segment_index": s["seg_index"],
                "chainage_start_m": s.get("chainage_start_m"),
                "chainage_end_m": s.get("chainage_end_m"),
                "chainage_mid_m": s.get("chainage_mid_m"),
                "approx_length_m": seg_len,
                "length_m": s.get("length_m"),
                "lon": s.get("lon"),
                "lat": s.get("lat"),
                "final_risk_score": s["final_score"],
                "risk_category_legacy": s["category"],
                "risk_category_0_100": s["risk_category_0_100"],
                "band": s["band"],
                "relative_risk_score": s["relative_risk_score"],
                "rank": s.get("rank"),
                "dominant_peril": s["dominant_peril"],
                "elev_m": s.get("elev_m"),
                "slope_deg": s.get("slope_deg"),
                "dist_water_proxy_m": s.get("dist_water_proxy_m"),
                "forecast_rain_mm": s.get("forecast_rain_mm"),
                "drv_precip_scaled": s.get("drv_precip_scaled"),
                "drv_dist_water_scaled": s.get("drv_dist_water_scaled"),
                "drv_lulc_scaled": s.get("drv_lulc_scaled"),
                "drv_slope_scaled": s.get("drv_slope_scaled"),
                "drv_dem_scaled": s.get("drv_dem_scaled"),
                "contrib_precip": s.get("contrib_precip"),
                "contrib_dist_water": s.get("contrib_dist_water"),
                "contrib_lulc": s.get("contrib_lulc"),
                "contrib_slope": s.get("contrib_slope"),
                "contrib_dem": s.get("contrib_dem"),
                "primary_driver": s.get("primary_driver"),
                "top_drivers": s.get("top_drivers"),
                "depth_100_undef_m": None,
                "eal_present": None,
                "result_state": "Modelled susceptibility (no calibrated AEP or depth)",
                "follow_up": "Inspect drainage and embankment if High or Extreme band",
            }
        )
    path = out_dir / f"Road_Section_Exposure_Register_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    print(f"  Wrote {path}")
    return path


def export_asset_scores(scored: List[Dict], out_dir: Path) -> Path:
    rows = [
        {
            "rank": s.get("rank"),
            "asset_id": s.get("asset_id"),
            "asset_name": s.get("asset_name"),
            "lat": s.get("lat"),
            "lon": s.get("lon"),
            "asset_type": s.get("asset_type"),
            "chainage_start_m": s.get("chainage_start_m"),
            "chainage_end_m": s.get("chainage_end_m"),
            "chainage_mid_m": s.get("chainage_mid_m"),
            "elev_m": s.get("elev_m"),
            "slope_deg": s.get("slope_deg"),
            "dist_water_proxy_m": s.get("dist_water_proxy_m"),
            "relative_risk_score": s["relative_risk_score"],
            "risk_category": s["risk_category_0_100"],
            "band": s["band"],
            "dominant_peril": s["dominant_peril"],
            "susceptibility_score_1_10": s["final_score"],
            "primary_driver": s.get("primary_driver"),
            "top_drivers": s.get("top_drivers"),
            "contrib_precip": s.get("contrib_precip"),
            "contrib_dist_water": s.get("contrib_dist_water"),
            "contrib_lulc": s.get("contrib_lulc"),
            "contrib_slope": s.get("contrib_slope"),
            "contrib_dem": s.get("contrib_dem"),
            "depth_100_undef": None,
            "eal_present": None,
            "horizon": "present",
            "scenario": "baseline",
        }
        for s in scored
    ]
    path = out_dir / f"Rozvi_Asset_Scores_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    print(f"  Wrote {path}")
    return path


def export_payload(
    scored: List[Dict],
    metrics: Dict[str, Any],
    layers: Dict[str, Any],
    seg_len: float,
    out_dir: Path,
    portfolio_name: str,
) -> Path:
    """JSON payload for downstream systems and template fill."""
    bc = metrics["band_counts"]
    n = metrics["asset_count"]
    payload = {
        "model_name": MODEL_NAME,
        "model_version": MODEL_VERSION,
        "portfolio_name": portfolio_name,
        "run_date": datetime.now().strftime("%Y-%m-%d"),
        "asset_count": n,
        "road_length_km": metrics["road_length_km"],
        "segment_length_m": seg_len,
        "dem_source": layers.get("dem_source"),
        "crs": "EPSG:4326",
        "driver_weights": layers.get("weights", DEFAULT_WEIGHTS),
        "driver_catalogue": layers.get("driver_catalogue", DRIVER_CATALOGUE),
        "climate_context": CLIMATE_CONTEXT,
        "model_quality": MODEL_QUALITY,
        "portfolio_mean_category": metrics["portfolio_mean_category"],
        "pct_assets_exposed_screening": metrics["pct_assets_exposed_screening"],
        "top_asset_name": metrics["top_asset_name"],
        "top_asset_category_score": metrics["top_asset_category_score"],
        "dominant_peril": metrics["dominant_peril"],
        "band_counts": bc,
        "assets": [
            {
                "asset_id": s["asset_id"],
                "section_id": s.get("section_id"),
                "lat": s["lat"],
                "lon": s["lon"],
                "chainage_start_m": s.get("chainage_start_m"),
                "chainage_end_m": s.get("chainage_end_m"),
                "elev_m": s.get("elev_m"),
                "slope_deg": s.get("slope_deg"),
                "susceptibility_1_10": s["final_score"],
                "risk_category": s["risk_category_0_100"],
                "band": s["band"],
                "relative_risk_score": s["relative_risk_score"],
                "rank": s.get("rank"),
                "primary_driver": s.get("primary_driver"),
                "top_drivers": s.get("top_drivers"),
                "contributions": {
                    "precip": s.get("contrib_precip"),
                    "dist_water": s.get("contrib_dist_water"),
                    "lulc": s.get("contrib_lulc"),
                    "slope": s.get("contrib_slope"),
                    "dem": s.get("contrib_dem"),
                },
            }
            for s in scored
        ],
    }
    path = out_dir / f"Rozvi_Report_Payload_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"  Wrote {path}")
    return path


def export_geojson(scored: List[Dict], out_dir: Path) -> Optional[Path]:
    try:
        gdf = gpd.GeoDataFrame(scored, geometry="geometry", crs="EPSG:4326")
        path = out_dir / f"Road_Segments_Scored_{datetime.now().strftime('%Y%m%d_%H%M%S')}.geojson"
        gdf.to_file(path, driver="GeoJSON")
        print(f"  Wrote {path}")
        return path
    except Exception as exc:
        print(f"  [warn] GeoJSON export failed: {exc}")
        return None


def _shade(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


def _hdr(row, color: str = "1F4E79") -> None:
    for cell in row.cells:
        _shade(cell, color)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
                r.font.name = "Arial"


def _font_row(row, size: int = 9) -> None:
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(size)


def generate_docx_report(
    scored: List[Dict],
    metrics: Dict[str, Any],
    layers: Dict[str, Any],
    seg_len: float,
    out_dir: Path,
    portfolio_name: str,
    figures: Optional[Dict[str, Path]] = None,
) -> Optional[Path]:
    if not HAS_DOCX:
        print("  [info] python-docx not installed; skipping Word report.")
        return None

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    def heading(text: str, level: int = 1) -> None:
        h = doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.name = "Arial"
            r.font.color.rgb = RGBColor(31, 78, 121)

    def para(text: str, size: int = 10, italic: bool = False) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.italic = italic
        p.paragraph_format.space_after = Pt(6)

    n = metrics["asset_count"]
    bc = metrics["band_counts"]
    catalogue = layers.get("driver_catalogue", DRIVER_CATALOGUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(MODEL_NAME)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)
    run.font.name = "Arial"
    para(portfolio_name, size=12)
    para(
        "Screening report. Flood depth, EAL and return-period hazard grids are not "
        "assessed in this build. Climate pathways are reported as regional context only.",
        size=9,
        italic=True,
    )

    # 0 Metadata
    heading("0. Report metadata")
    meta = [
        ("Report title", f"Flood Risk Report – {portfolio_name}"),
        ("Model", f"{MODEL_NAME} {MODEL_VERSION}"),
        ("DEM source", str(layers.get("dem_source"))),
        ("Run date", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Assets assessed", str(n)),
        ("Road length (km)", str(metrics["road_length_km"])),
        ("Segment length (m)", str(seg_len)),
        ("CRS", "EPSG:4326"),
        ("Calibration status", MODEL_QUALITY["calibration_status"]),
    ]
    tbl = doc.add_table(rows=len(meta), cols=2)
    tbl.style = "Table Grid"
    for i, (a, b) in enumerate(meta):
        tbl.rows[i].cells[0].text = a
        tbl.rows[i].cells[1].text = b
        _font_row(tbl.rows[i])
        tbl.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        _shade(tbl.rows[i].cells[0], "D6E3F0")

    # 1 Executive summary
    heading("1. Executive summary")
    heading("1.1 Headline metrics", 2)
    head = [
        ("Metric", "Value"),
        (
            f"Assets at screening threshold (score >= {SCREENING_THRESHOLD})",
            f"{metrics['n_assets_exposed_screening']} of {n} ({metrics['pct_assets_exposed_screening']}%)",
        ),
        (
            "Highest-risk segment",
            f"{metrics['top_asset_name']} (category {metrics['top_asset_category_score']})",
        ),
        ("Mean risk category (0-100)", str(metrics["portfolio_mean_category"])),
        ("Dominant peril", metrics["dominant_peril"]),
        ("EAL / 1-in-100 depth", "not assessed"),
    ]
    ht = doc.add_table(rows=len(head), cols=2)
    ht.style = "Table Grid"
    for i, (a, b) in enumerate(head):
        ht.rows[i].cells[0].text = a
        ht.rows[i].cells[1].text = b
        _font_row(ht.rows[i])
    _hdr(ht.rows[0])

    heading("1.2 Risk distribution", 2)
    dist = [("Band", "Range", "Count", "%")]
    for name, lo, hi in RISK_BANDS:
        c = bc.get(name, 0)
        dist.append((name, f"{lo}–{hi}", str(c), f"{round(c / n * 100, 1) if n else 0}%"))
    dt = doc.add_table(rows=len(dist), cols=4)
    dt.style = "Table Grid"
    for i, row in enumerate(dist):
        for j, val in enumerate(row):
            dt.rows[i].cells[j].text = val
        _font_row(dt.rows[i])
    _hdr(dt.rows[0])

    figures = figures or {}
    if figures.get("band_distribution"):
        para("Figure 1. Risk band distribution across assessed segments.", size=9, italic=True)
        doc.add_picture(str(figures["band_distribution"]), width=Inches(5.8))
    if figures.get("score_histogram"):
        para("Figure 2. Distribution of susceptibility scores (1–10).", size=9, italic=True)
        doc.add_picture(str(figures["score_histogram"]), width=Inches(5.8))

    # 2 Method + drivers
    heading("2. Method")
    para(
        "Susceptibility is a weighted combination of the drivers listed below. "
        "Each driver is scaled to the interval 1-10 using AOI percentile stretches "
        "(2nd-98th). Segment scores are mean values sampled under a corridor buffer."
    )
    para(f"DEM source: {layers.get('dem_source')}.")

    heading("2.1 Flood risk drivers and weights", 2)
    drv_header = ["Driver", "Weight", "Weight %", "Direction", "Data used in this run"]
    drv = doc.add_table(rows=1 + len(catalogue), cols=len(drv_header))
    drv.style = "Table Grid"
    for j, c in enumerate(drv_header):
        drv.rows[0].cells[j].text = c
    _hdr(drv.rows[0])
    for i, d in enumerate(catalogue):
        vals = [
            d["name"],
            f"{d['weight']:.2f}",
            f"{d['weight'] * 100:.0f}%",
            d["direction"],
            str(d.get("data_status", "")),
        ]
        for j, v in enumerate(vals):
            drv.rows[i + 1].cells[j].text = v
        _font_row(drv.rows[i + 1], size=8)
    para(
        f"Weights sum to {sum(d['weight'] for d in catalogue):.2f}. "
        "Land cover is held at a neutral scaled value when no land-cover raster is supplied.",
        size=9,
        italic=True,
    )

    heading("2.2 Climate context (CMIP6 / CMIP7)", 2)
    para(f"Framework referenced: {CLIMATE_CONTEXT['framework']}.")
    para(f"Pathways cited for narrative context: {CLIMATE_CONTEXT['pathways_referenced']}.")
    para(CLIMATE_CONTEXT["application"])
    para(CLIMATE_CONTEXT["cmip7_note"], size=9, italic=True)
    para(
        "Segment-scale climate-adjusted depths and probabilities: "
        f"{CLIMATE_CONTEXT['segment_scale_climate']}.",
        size=9,
    )

    heading("2.3 Model quality", 2)
    quality_rows = [
        ("Purpose", MODEL_QUALITY["purpose"]),
        ("Calibration", MODEL_QUALITY["calibration_status"]),
        ("Validation", MODEL_QUALITY["validation_status"]),
        ("Spatial unit", MODEL_QUALITY["spatial_unit"]),
        ("Vertical data limit", MODEL_QUALITY["vertical_data_limit"]),
        ("Missing data policy", MODEL_QUALITY["missing_data_policy"]),
        ("Suitable uses", MODEL_QUALITY["suitable_uses"]),
        ("Unsuitable uses", MODEL_QUALITY["unsuitable_uses"]),
    ]
    qt = doc.add_table(rows=len(quality_rows), cols=2)
    qt.style = "Table Grid"
    for i, (a, b) in enumerate(quality_rows):
        qt.rows[i].cells[0].text = a
        qt.rows[i].cells[1].text = b
        _font_row(qt.rows[i], size=8)
        qt.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        _shade(qt.rows[i].cells[0], "D6E3F0")

    if figures.get("corridor_map"):
        heading("2.4 Corridor map", 2)
        para(
            "Segments coloured by risk band. This is a susceptibility ranking map, "
            "not a flood inundation extent product.",
            size=9,
            italic=True,
        )
        doc.add_picture(str(figures["corridor_map"]), width=Inches(5.8))

    if figures.get("longitudinal_profile"):
        heading("2.5 Longitudinal risk profile", 2)
        para(
            "Susceptibility score along the assessed network in alignment order. "
            "The dashed line marks the screening threshold.",
            size=9,
            italic=True,
        )
        doc.add_picture(str(figures["longitudinal_profile"]), width=Inches(6.2))

    if figures.get("primary_driver_share"):
        heading("2.6 Primary driver across the portfolio", 2)
        para(
            "Count of segments by the driver with the largest weighted contribution "
            "at that location.",
            size=9,
            italic=True,
        )
        doc.add_picture(str(figures["primary_driver_share"]), width=Inches(5.8))

    # 3 Priority segments
    heading("3. Priority segments")
    para(
        "For each segment the composite score uses the fixed driver weights. "
        "Primary driver and top drivers are the largest weighted contributions "
        "at that location (contribution = weight × scaled driver value).",
        size=9,
        italic=True,
    )
    top15 = sorted(scored, key=lambda x: x["risk_category_0_100"], reverse=True)[:15]
    cols = [
        "Rank", "Section", "Chainage (m)", "Score", "Cat", "Band",
        "Primary driver", "Top drivers (contribution)",
    ]
    mt = doc.add_table(rows=1 + len(top15), cols=len(cols))
    mt.style = "Table Grid"
    for j, c in enumerate(cols):
        mt.rows[0].cells[j].text = c
    _hdr(mt.rows[0])
    for i, s in enumerate(top15):
        vals = [
            str(s.get("rank")),
            str(s.get("section_id")),
            f"{s.get('chainage_start_m')}–{s.get('chainage_end_m')}",
            str(s.get("final_score")),
            str(s.get("risk_category_0_100")),
            str(s.get("band")),
            str(s.get("primary_driver")),
            str(s.get("top_drivers")),
        ]
        for j, v in enumerate(vals):
            mt.rows[i + 1].cells[j].text = v
        _font_row(mt.rows[i + 1], size=8)

    if figures.get("driver_contributions"):
        para(
            "Figure. Weighted driver contributions for the highest-ranked segments. "
            "Bar length is the contribution to the composite score (weight × scaled driver).",
            size=9,
            italic=True,
        )
        doc.add_picture(str(figures["driver_contributions"]), width=Inches(6.2))

    heading("4. Segment and chainage fields")
    para(
        f"Network divided into {n} segments of nominal length {seg_len} m. "
        "Chainage is sequential from the first vertex of each parent road feature "
        "and should be replaced with official alignment chainage when available."
    )

    heading("5. Limitations")
    for line in [
        "Outputs are relative susceptibility rankings, not calibrated probabilities or depths.",
        "No return-period depth grids, defence performance, or segment-scale climate depth adjustment.",
        "Coarse DEM sources cannot resolve low embankments or bridge decks.",
        "Chainage origin is model-derived unless replaced with official alignment data.",
        "Missing raster samples default toward a neutral susceptibility contribution.",
        "CMIP pathways are cited for regional context only; they do not alter segment scores in this version.",
    ]:
        para(f"• {line}")

    path = out_dir / f"Flood_Risk_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(str(path))
    print(f"  Wrote {path}")
    return path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description=f"{MODEL_NAME} – corridor flood susceptibility screening"
    )
    p.add_argument("--roads", required=True, help="Road network GeoJSON or Shapefile")
    p.add_argument("--dem", default=None, help="Optional DEM GeoTIFF (recommended)")
    p.add_argument("--water", default=None, help="Optional water-mask GeoTIFF")
    p.add_argument("--rain", default=None, help="Optional rainfall GeoTIFF (mm)")
    p.add_argument("--seg-length", type=float, default=SEGMENT_LENGTH_M)
    p.add_argument("--corridor", type=float, default=CORRIDOR_HALF_WIDTH_M)
    p.add_argument("--max-segments", type=int, default=MAX_SEGMENTS)
    p.add_argument("--outdir", default="./outputs")
    p.add_argument("--portfolio-name", default="Road corridor portfolio")
    p.add_argument("--no-docx", action="store_true", help="Skip Word report")
    return p.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    out_dir = Path(args.outdir)
    out_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 70)
    print(f"{MODEL_NAME}  v{MODEL_VERSION}")
    print("=" * 70)

    roads_path = Path(args.roads)
    if not roads_path.exists():
        print(f"ERROR: road file not found: {roads_path}")
        return 1

    print(f"\n[1/7] Loading roads: {roads_path}")
    roads = ensure_wgs84(gpd.read_file(roads_path))
    roads = roads[roads.geometry.type.isin(["LineString", "MultiLineString"])].copy()
    if roads.empty:
        print("ERROR: no LineString / MultiLineString features.")
        return 1
    print(f"  {len(roads)} line feature(s)")

    print(f"\n[2/7] Segmenting (~{args.seg_length:.0f} m) ...")
    segments = segment_roads(roads, seg_len_m=args.seg_length, max_segments=args.max_segments)
    print(f"  {len(segments)} segments")

    minx, miny, maxx, maxy = segments.total_bounds
    pad = 0.02
    bounds = (minx - pad, miny - pad, maxx + pad, maxy + pad)
    print(f"\n[3/7] Building risk layers {tuple(round(b, 4) for b in bounds)}")
    layers = build_risk_layers(
        bounds,
        dem_path=Path(args.dem) if args.dem else None,
        water_path=Path(args.water) if args.water else None,
        rain_path=Path(args.rain) if args.rain else None,
    )
    print("  Driver weights:", layers["weights"])

    print(f"\n[4/7] Scoring segments ...")
    scored = score_segments(segments, layers, corridor_half_width_m=args.corridor)
    scored = assign_chainage(scored, args.seg_length)
    metrics = portfolio_metrics(scored, args.seg_length)
    print("  Band counts:", metrics["band_counts"])

    print(f"\n[5/7] Writing outputs -> {out_dir.resolve()}")
    export_driver_register(layers, out_dir)
    export_executive_summary(scored, args.seg_length, metrics, layers, out_dir)
    export_section_register(scored, args.seg_length, out_dir)
    export_asset_scores(scored, out_dir)
    export_payload(scored, metrics, layers, args.seg_length, out_dir, args.portfolio_name)
    export_geojson(scored, out_dir)

    print("\n[6/7] Figures")
    figures = generate_figures(scored, metrics, out_dir)

    print("\n[7/7] Report")
    if not args.no_docx:
        generate_docx_report(
            scored,
            metrics,
            layers,
            args.seg_length,
            out_dir,
            args.portfolio_name,
            figures=figures,
        )

    print("\nComplete.")
    print(
        "Note: results are susceptibility rankings for screening. "
        "They are not calibrated flood probabilities or water-depth estimates."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())